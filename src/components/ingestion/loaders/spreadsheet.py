from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from ..base import BaseLoader
from ..exceptions import OptionalDependencyError
from ..models import DocumentBlock


class SpreadsheetLoader(BaseLoader):
    name = "spreadsheet-openpyxl"
    extensions = (".xlsx", ".xlsm", ".xltx", ".xltm")

    def load_blocks(self, path: Path) -> Iterable[DocumentBlock]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise OptionalDependencyError(
                "Excel ingestion requires `openpyxl`."
            ) from exc

        wb = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
        )

        try:
            for ws in wb.worksheets:
                rows = []
                for row_idx, row in enumerate(
                    ws.iter_rows(
                        max_row=self.config.spreadsheet_max_rows,
                        max_col=self.config.spreadsheet_max_columns,
                        values_only=True,
                    ),
                    start=1,
                ):
                    values = ["" if value is None else str(value) for value in row]

                    if not any(v.strip() for v in values):
                        continue

                    rows.append(" | ".join(values))

                if rows:
                    yield DocumentBlock(
                        block_id="",
                        text="\n".join(rows),
                        block_type="table",
                        sheet=ws.title,
                        source_locator=f"sheet:{ws.title}",
                        metadata={"sheet": ws.title},
                    )
        finally:
            wb.close()


class CSVLoader(BaseLoader):
    name = "csv-stdlib"
    extensions = (".csv", ".tsv")

    def load_blocks(self, path: Path) -> Iterable[DocumentBlock]:
        raw = path.read_bytes()
        text = None

        for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                pass

        text = text or raw.decode("utf-8", errors="replace")

        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel

        reader = csv.reader(text.splitlines(), dialect)
        rows = []

        for row_index, row in enumerate(reader):
            if row_index >= self.config.spreadsheet_max_rows:
                break
            rows.append(" | ".join(cell.strip() for cell in row))

        if rows:
            yield DocumentBlock(
                block_id="",
                text="\n".join(rows),
                block_type="table",
                source_locator=str(path.name),
            )

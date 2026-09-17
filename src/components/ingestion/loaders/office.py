from __future__ import annotations

from pathlib import Path
from typing import Iterable

from ..base import BaseLoader
from ..exceptions import OptionalDependencyError
from ..models import DocumentBlock


class DOCXLoader(BaseLoader):
    name = "docx-python-docx"
    extensions = (".docx",)

    def load_blocks(self, path: Path) -> Iterable[DocumentBlock]:
        try:
            from docx import Document as WordDocument
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise OptionalDependencyError(
                "DOCX ingestion requires `python-docx`."
            ) from exc

        document = WordDocument(path)

        # Iterate body children so paragraph/table order is preserved.
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                if text:
                    style = paragraph.style.name if paragraph.style else None
                    block_type = "heading" if style and "Heading" in style else "paragraph"
                    yield DocumentBlock(
                        block_id="",
                        text=text,
                        block_type=block_type,
                        section=style,
                        source_locator=f"paragraph",
                        metadata={"style": style},
                    )

            elif child.tag.endswith("}tbl"):
                table = Table(child, document)
                rows = []
                for row in table.rows:
                    rows.append(" | ".join(cell.text.strip() for cell in row.cells))
                text = "\n".join(rows).strip()
                if text:
                    yield DocumentBlock(
                        block_id="",
                        text=text,
                        block_type="table",
                        source_locator="table",
                    )


class PPTXLoader(BaseLoader):
    name = "pptx-python-pptx"
    extensions = (".pptx",)

    def load_blocks(self, path: Path) -> Iterable[DocumentBlock]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise OptionalDependencyError(
                "PPTX ingestion requires `python-pptx`."
            ) from exc

        presentation = Presentation(path)

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_parts: list[str] = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_parts.append(
                            " | ".join(cell.text.strip() for cell in row.cells)
                        )

            text = "\n".join(slide_parts).strip()
            if text:
                yield DocumentBlock(
                    block_id="",
                    text=text,
                    block_type="slide",
                    slide=slide_number,
                    source_locator=f"slide:{slide_number}",
                )


class ODTLoader(BaseLoader):
    name = "odt-odfpy"
    extensions = (".odt",)

    def load_blocks(self, path: Path) -> Iterable[DocumentBlock]:
        try:
            from odf import text, teletype
            from odf.opendocument import load
        except ImportError as exc:
            raise OptionalDependencyError(
                "ODT ingestion requires `odfpy`."
            ) from exc

        document = load(path)

        for element in document.getElementsByType(text.P):
            value = teletype.extractText(element).strip()
            if value:
                yield DocumentBlock(
                    block_id="",
                    text=value,
                    block_type="paragraph",
                    source_locator="paragraph",
                )

        for element in document.getElementsByType(text.H):
            value = teletype.extractText(element).strip()
            if value:
                yield DocumentBlock(
                    block_id="",
                    text=value,
                    block_type="heading",
                    source_locator="heading",
                )
